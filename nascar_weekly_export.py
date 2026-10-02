#!/usr/bin/env python3
"""
Weekly NASCAR export: two tables per race, as CSV, Excel and PNG.

  Table 1  Race results   (finish, start, stage pts, race pts, laps led, points after race)
  Table 2  Season-to-date (rank, +/-, points, behind, stats, finish in recent races)

Usage
  python nascar_weekly_export.py results_2026.csv --week 30 --outdir out
  python nascar_weekly_export.py results_2026.csv --all-weeks --outdir out
  Flags: --no-png  --no-xlsx  --png-races 5  --title "2026 NASCAR Cup Series Chase Scorecard"
  Fastest laps and race events are picked up automatically when fastest_laps.csv and
  race_events.csv sit next to the results file (turn off with --no-auto), and they are added to
  each week's Excel file as "Fastest laps" and "Race events" sheets and to the race image.
  Race bar (stages / yellows / red flag length above the race results image):
    --events race_events.csv   columns race_num, kind, start_lap, end_lap, minutes
        kind = total (end_lap = race laps) | stage (end_lap) | caution (start_lap, end_lap)
             | red (start_lap = lap of the stoppage, minutes = length)
  Fastest laps (not in the nascaR.data file, so supplied separately):
    --fastest-laps fastest_laps.csv   built by fetch_best_laps.py; columns race_num, driver,
                                      best_lap_time (seconds), optional best_lap_speed (mph),
                                      best_lap_num, fast_laps (laps run fastest of the field)

Chase reset (2026 format, read from the xfile345 scorecard -- confirm against the rules):
  After race 26 the top 16 in points are reset to 2000 + a regular-season bonus
  (100, 75, 65, then 60 down to 0 in steps of 5 by regular-season rank). Drivers outside the
  top 16 keep their season totals and rank below the Chase field, with "behind" measured
  from their own group's leader. Change with --reset-after / --chase-size / --reset-base,
  or turn off with --no-reset.  The reset is added as adjustments in race 27.

results.csv  (one row per driver per race)
  race_num, driver, finish, race_pts            required
  race_name, track, car, start, stage1_pts, stage2_pts, laps_led, dnf, pole, stage_wins   optional

adj.csv  (optional: penalties, bonuses, anything the race-by-race sum misses)
  race_num, driver, points, note      applied from race_num onward
"""
import argparse
import os
import re
import sys

import pandas as pd

REQUIRED = ["race_num", "driver", "finish", "race_pts"]
DEFAULTS = {
    "race_name": "", "track": "", "car": "", "start": pd.NA, "stage1_pts": 0, "stage2_pts": 0,
    "laps_led": 0, "dnf": 0, "pole": 0, "stage_wins": 0,
    "led_most": 0, "s1_win": 0, "s2_win": 0, "rating": float("nan"),
}
# regular-season bonus by regular-season rank (from the 2026 scorecard image)
BONUSES = [100, 75, 65] + list(range(60, -1, -5))


# ----------------------------------------------------------------------------- loading
def load_results(path):
    df = pd.read_csv(path, dtype={"car": str})       # keep "08" distinct from "8"
    missing = [c for c in REQUIRED if c not in df.columns]
    if missing:
        sys.exit(f"results file is missing columns: {missing}")
    for col, val in DEFAULTS.items():
        if col not in df.columns:
            df[col] = val
    df["car"] = df["car"].astype(str).replace("<NA>", "").replace("nan", "")
    df["rating"] = pd.to_numeric(df["rating"], errors="coerce")
    derived = []
    if (df["led_most"] == 0).all():
        most = df.groupby("race_num")["laps_led"].transform("max")
        df["led_most"] = ((df["laps_led"] > 0) & (df["laps_led"] == most)).astype(int)
        derived.append("led_most")
    if (df["s1_win"] == 0).all() and (df["s2_win"] == 0).all():
        df["s1_win"] = (pd.to_numeric(df["stage1_pts"], errors="coerce") == 10).astype(int)
        df["s2_win"] = (pd.to_numeric(df["stage2_pts"], errors="coerce") == 10).astype(int)
        derived.append("stage wins")
    if (df["pole"] == 0).all() and df["start"].notna().any():
        df["pole"] = (pd.to_numeric(df["start"], errors="coerce") == 1).astype(int)
        derived.append("pole (start = 1)")
    if (df["stage_wins"] == 0).all():
        df["stage_wins"] = df["s1_win"] + df["s2_win"]
    if derived:
        print("note: results file had no flags for " + ", ".join(derived) +
              "; worked them out from the other columns (rerun the import for exact flags)")
    return df


def load_adjustments(path):
    if not path:
        return pd.DataFrame(columns=["race_num", "driver", "points", "note"])
    adj = pd.read_csv(path)
    if "note" not in adj.columns:
        adj["note"] = ""
    return adj


def _norm(name):
    """Loose driver-name key so 'Ricky Stenhouse Jr.' matches 'Ricky Stenhouse Jr' etc."""
    n = re.sub(r"\(.*?\)", "", str(name).lower()).replace("-", " ")
    n = re.sub(r"[^a-z0-9 ]", "", n)
    parts = [w for w in n.split() if w not in ("jr", "sr", "ii", "iii")]
    return " ".join(parts)


def merge_fastest(results, path):
    """Merge per-driver best laps: race_num, driver, best_lap_time (seconds)
    [, best_lap_speed (mph), best_lap_num]. Adds fl_race (1 = fastest lap of that race)."""
    fl = pd.read_csv(path)
    need = ["race_num", "driver", "best_lap_time"]
    if any(c not in fl.columns for c in need):
        sys.exit(f"{path} needs columns {need} (optional: best_lap_speed, best_lap_num)")
    for c in ("best_lap_speed", "best_lap_num", "fast_laps"):
        if c not in fl.columns:
            fl[c] = float("nan")
    for c in ("best_lap_time", "best_lap_speed", "best_lap_num", "fast_laps"):
        fl[c] = pd.to_numeric(fl[c], errors="coerce")
    results = results.copy()
    results["_k"] = results["driver"].map(_norm)
    fl["_k"] = fl["driver"].map(_norm)
    known = set(zip(results["race_num"], results["_k"]))
    lost = fl[[(r, k) not in known for r, k in zip(fl["race_num"], fl["_k"])]]
    if len(lost):
        print(f"note: {len(lost)} fastest-lap rows matched no driver in the results "
              f"(e.g. {', '.join(lost['driver'].astype(str).head(4))})")
    m = fl.drop_duplicates(["race_num", "_k"]).set_index(["race_num", "_k"])[
        ["best_lap_time", "best_lap_speed", "best_lap_num", "fast_laps"]]
    results = results.join(m, on=["race_num", "_k"]).drop(columns="_k")
    fastest = results.groupby("race_num")["best_lap_time"].transform("min")
    results["fl_race"] = (results["best_lap_time"].notna() & (results["best_lap_time"] == fastest)).astype(int)
    results["fl_rank"] = results.groupby("race_num")["best_lap_time"].rank(method="min")
    return results


def _to_minutes(v):
    """42 -> 42.0 ; '42:10' (mm:ss) -> 42.17 ; '1:14:00' (h:mm:ss) -> 74.0 ; blank -> NaN"""
    if v is None or (isinstance(v, float) and pd.isna(v)) or str(v).strip() == "":
        return float("nan")
    t = str(v).strip()
    if ":" not in t:
        return float(t)
    parts = [float(x) for x in t.split(":")]
    if len(parts) == 3:
        return parts[0] * 60 + parts[1] + parts[2] / 60
    if len(parts) == 2:
        return parts[0] + parts[1] / 60
    raise ValueError(f"can't read duration '{v}' (use minutes, mm:ss or h:mm:ss)")


def load_events(path):
    """race_events.csv: race_num, kind, start_lap, end_lap, minutes, note
       kind = stage (end_lap = last lap of the stage) | total (end_lap = race length)
            | caution (start_lap, end_lap) | red (start_lap = lap of the stoppage, minutes = length)"""
    ev = pd.read_csv(path)
    need = ["race_num", "kind"]
    if any(c not in ev.columns for c in need):
        sys.exit(f"{path} needs columns race_num and kind (plus start_lap, end_lap, minutes)")
    for alias in ("duration", "time"):                 # accept either header for the red-flag length
        if "minutes" not in ev.columns and alias in ev.columns:
            ev["minutes"] = ev[alias]
    for c in ("start_lap", "end_lap", "minutes"):
        if c not in ev.columns:
            ev[c] = float("nan")
    ev["minutes"] = ev["minutes"].map(_to_minutes)
    ev["kind"] = ev["kind"].str.lower().str.strip()
    return ev


def bar_for_race(events, week):
    """Turn the event rows for one race into the dict the bar drawing uses (None if unusable)."""
    if events is None:
        return None
    e = events[events["race_num"] == week]
    tot = e[e["kind"] == "total"]["end_lap"].dropna()
    if tot.empty:
        return None
    return dict(
        total=float(tot.iloc[0]),
        stages=sorted(float(x) for x in e[e["kind"] == "stage"]["end_lap"].dropna()),
        cautions=[(float(a), float(b)) for a, b in
                  e[e["kind"] == "caution"][["start_lap", "end_lap"]].dropna().itertuples(index=False)],
        reds=[(float(l), None if pd.isna(m) else float(m)) for l, m in
              e[e["kind"] == "red"][["start_lap", "minutes"]].dropna(subset=["start_lap"]).itertuples(index=False)],
    )


def events_table(events, week):
    """One race's stages / cautions / red flags as a readable table (None if there are none)."""
    if events is None:
        return None
    e = events[events["race_num"] == week]
    if e.empty:
        return None
    rows = []
    for i, (_, r) in enumerate(e[e["kind"] == "stage"].sort_values("end_lap").iterrows(), 1):
        rows.append(dict(_k=r["end_lap"], Event=f"Stage {i} ends", **{"End Lap": r["end_lap"]}))
    for _, r in e[e["kind"] == "caution"].sort_values("start_lap").iterrows():
        rows.append(dict(_k=r["start_lap"], Event="Caution", **{"Start Lap": r["start_lap"], "End Lap": r["end_lap"],
                         "Laps": r["end_lap"] - r["start_lap"] + 1}))
    for _, r in e[e["kind"] == "red"].sort_values("start_lap").iterrows():
        rows.append(dict(_k=r["start_lap"], Event="Red flag", **{"Start Lap": r["start_lap"],
                         "Duration": _dur(None if pd.isna(r["minutes"]) else r["minutes"]), "Note": r.get("note", "")}))
    for _, r in e[e["kind"] == "total"].iterrows():
        rows.append(dict(_k=1e9, Event="Race length", Laps=r["end_lap"]))
    t = pd.DataFrame(rows).sort_values("_k", kind="stable").drop(columns="_k")
    for c in ("Start Lap", "End Lap", "Laps"):
        if c in t:
            t[c] = t[c].astype("Int64")
    return t[[c for c in ("Event", "Start Lap", "End Lap", "Laps", "Duration", "Note") if c in t]]


def fastest_table(results, week):
    """One race's best laps, quickest first (None if no lap data was merged)."""
    if "best_lap_time" not in results:
        return None
    r = results[(results.race_num == week) & results.best_lap_time.notna()].sort_values("best_lap_time")
    if r.empty:
        return None
    t = pd.DataFrame({"Rank": r["fl_rank"].astype("Int64"), "Car": r["car"], "Driver": r["driver"],
                      "Best Lap": r["best_lap_time"], "Mph": r["best_lap_speed"].round(2),
                      "Lap #": r["best_lap_num"].astype("Int64"), "Fast Laps": r["fast_laps"].astype("Int64")})
    return t.dropna(axis=1, how="all").reset_index(drop=True)


# ----------------------------------------------------------------------------- points
def points_through(results, adj, week):
    """Total points per driver after `week` = sum of race_pts + adjustments."""
    r = results[results.race_num <= week].groupby("driver")["race_pts"].sum()
    a = adj[adj.race_num <= week].groupby("driver")["points"].sum()
    return r.add(a, fill_value=0).astype(int)


def make_chase(results, adj, after, size, base):
    """Return (adjustments incl. the reset, chase info). chase is None if no reset applies."""
    if after is None or results.race_num.max() <= after:
        return adj, None
    pts = points_through(results, adj, after)
    up = results[results.race_num <= after]
    wins = up[up.finish == 1].groupby("driver").size().reindex(pts.index, fill_value=0)
    avg = up.groupby("driver")["finish"].mean().reindex(pts.index)
    order = pd.DataFrame({"pts": pts, "wins": wins, "avg": avg}).sort_values(
        ["pts", "wins", "avg"], ascending=[False, False, True], kind="stable")
    field = list(order.index[:size])
    reg_rank = pd.Series(range(1, len(order) + 1), index=order.index)
    bonus = pd.Series((BONUSES + [0] * size)[:size], index=field)
    rows = [dict(race_num=after + 1, driver=d, points=base + int(bonus[d]) - int(pts[d]),
                 note="Chase reset") for d in field]
    adj2 = pd.concat([adj, pd.DataFrame(rows)], ignore_index=True)
    return adj2, dict(after=after, field=field, reg_rank=reg_rank, bonus=bonus)


def _grouped(week, chase):
    return bool(chase) and week > chase["after"]


def rank_series(pts, week, chase):
    if _grouped(week, chase):
        inch = pts.index.isin(chase["field"])
        a = pts[inch].rank(method="min", ascending=False)
        b = pts[~inch].rank(method="min", ascending=False) + inch.sum()
        return pd.concat([a, b]).astype(int)
    return pts.rank(method="min", ascending=False).astype(int)


def behind_series(pts, week, chase):
    if _grouped(week, chase):
        inch = pts.index.isin(chase["field"])
        a = pts[inch] - pts[inch].max()
        b = pts[~inch] - (pts[~inch].max() if (~inch).any() else 0)
        return pd.concat([a, b])
    return pts - pts.max()


# ----------------------------------------------------------------------------- ratings
def rating_stats(results, week, n=5):
    """Season-average rating and rolling rating (last n rated starts) per driver through `week`."""
    r = results[(results.race_num <= week)].dropna(subset=["rating"]).sort_values("race_num")
    g = r.groupby("driver")["rating"]
    return g.mean().round(1), g.apply(lambda s: s.tail(n).mean()).round(1)


# ----------------------------------------------------------------------------- tables
def standings_table(results, adj, week, chase=None):
    pts = points_through(results, adj, week)
    grouped = _grouped(week, chase)
    now_rank = rank_series(pts, week, chase)
    if week > 1:
        prev = points_through(results, adj, week - 1)
        prev = prev[prev.index.isin(pts.index)]
        prev_rank = rank_series(prev, week - 1, chase)

    upto = results[results.race_num <= week].copy()
    upto["top5"] = upto.finish <= 5
    upto["top10"] = upto.finish <= 10
    upto["top20"] = upto.finish <= 20
    upto["p30"] = upto.finish >= 30
    upto["win"] = upto.finish == 1
    upto["led_any"] = upto.laps_led > 0
    g = upto.groupby("driver")
    stats = pd.DataFrame({
        "Car": g["car"].last(),
        "Stage Pts": g["stage1_pts"].sum() + g["stage2_pts"].sum(),
        "Wins": g["win"].sum(), "Top 5": g["top5"].sum(), "Top 10": g["top10"].sum(),
        "Top 20": g["top20"].sum(), "30+": g["p30"].sum(), "DNFs": g["dnf"].sum(),
        "Poles": g["pole"].sum(), "Stage Wins": g["stage_wins"].sum(),
        "Races Led": g["led_any"].sum(), "Laps Led": g["laps_led"].sum(),
        "Races": g["finish"].count(), "Best Finish": g["finish"].min(),
        "Avg Finish": g["finish"].mean().round(1),
    })

    if "fl_race" in upto:
        stats["Best Lap Pts"] = g["fl_race"].sum()
    avg_rt, l5_rt = rating_stats(results, week)
    stats["Rating"], stats["Last 5 Rtg"] = avg_rt, l5_rt
    out = stats.join(pts.rename("Points"))
    out["Behind"] = behind_series(pts, week, chase)
    out["Rank"] = now_rank
    out["+/-"] = (prev_rank.reindex(out.index) - now_rank).fillna(0).astype(int) if week > 1 else 0

    if grouped:
        out["In Chase"] = ["Y" if d in chase["field"] else "" for d in out.index]
        out["Reg Rank"] = chase["reg_rank"].reindex(out.index).astype("Int64")
        out["Reg Bonus"] = chase["bonus"].reindex(out.index).astype("Int64")
        out["_g"] = (out["In Chase"] != "Y").astype(int)
    else:
        out["_g"] = 0

    grid = upto.pivot_table(index="driver", columns="race_num", values="finish", aggfunc="first")
    grid.columns = [f"R{c}" for c in grid.columns]
    out = out.join(grid)

    out = out.sort_values(["_g", "Points", "Wins"], ascending=[True, False, False], kind="stable")
    out = out.drop(columns="_g").reset_index().rename(columns={"driver": "Driver"})
    counts = out["Rank"].value_counts()
    out["Rank"] = out["Rank"].map(lambda r: f"T-{r}" if counts[r] > 1 else str(r))
    front = ["Rank", "+/-", "Car", "Driver", "Points", "Behind"]
    rest = [c for c in out.columns if c not in front]
    return out[front + rest]


def race_table(results, adj, week, chase=None):
    race = results[results.race_num == week].copy()
    pts = points_through(results, adj, week)
    rank = rank_series(pts, week, chase)
    _, l5 = rating_stats(results, week)
    race["Rating"] = race["rating"]
    race["Last 5 Rtg"] = race["driver"].map(l5)
    race["Total Pts"] = race["driver"].map(pts)
    race["Rank After"] = race["driver"].map(rank)
    race["Race Pts"] = race["race_pts"]
    race["Led"] = race["laps_led"]
    race = race.sort_values("finish")
    cols = {
        "finish": "Finish", "start": "Start", "car": "Car", "driver": "Driver",
        "stage1_pts": "Stg 1", "stage2_pts": "Stg 2", "Race Pts": "Race Pts",
        "Led": "Laps Led", "dnf": "DNF", "Rating": "Rating", "Last 5 Rtg": "Last 5 Rtg",
    }
    has_bl = "best_lap_time" in race and race["best_lap_time"].notna().any()
    if has_bl:
        race["best_lap_speed"] = race["best_lap_speed"].round(2) if "best_lap_speed" in race else float("nan")
        race["fl_rank"] = race["fl_rank"].astype("Int64")
        cols.update({"best_lap_time": "Best Lap", "fl_rank": "FL Rk"})
        if race["best_lap_speed"].notna().any():
            cols["best_lap_speed"] = "Mph"
        if "fast_laps" in race and race["fast_laps"].notna().any():
            race["fast_laps"] = race["fast_laps"].astype("Int64")
            cols["fast_laps"] = "Fast Laps"
    cols.update({"Total Pts": "Total Pts", "Rank After": "Rank After"})
    out = race[list(cols)].rename(columns=cols).reset_index(drop=True)
    return out.drop(columns=["Rating", "Last 5 Rtg"]) if out["Rating"].isna().all() else out


# ----------------------------------------------------------------------------- cell styling
# Mirrors the scorecard legend: italic = pole, bold = led a lap, red = led most laps,
# blue = won one stage, pink = won both stages, strikethrough = DNF, underline = fastest lap.
PINK, BLUE, RED = "#e0329e", "#1f55d6", "#d11a1a"
LEGEND = [
    ("Pole", dict(italic=True)),
    ("Led Lap", dict(bold=True)),
    ("Led Most", dict(bold=True, color=RED)),
    ("Won a Stage", dict(bold=True, color=BLUE)),
    ("Won Both Stages", dict(bold=True, color=PINK)),
    ("DNF", dict(strike=True)),
    ("Fastest Lap", dict(underline=True)),
]


def cell_styles(results, t2, cols, shade=True, shade_below=False):
    """{(row_index, column): style} for the race-finish cells in `cols` of the standings table."""
    idx = results.set_index(["race_num", "driver"])
    styles = {}
    for i, driver in enumerate(t2["Driver"]):
        for c in cols:
            key = (int(c[1:]), driver)
            if key not in idx.index:
                continue
            r = idx.loc[key]
            st = {}
            if shade:                                     # same tiers as the race results image
                f = r["finish"]
                if f <= 5:
                    st["fill"] = TOP5_FILL
                elif f <= 10:
                    st["fill"] = TOP10_FILL
                elif shade_below and f > 15:
                    st["fill"] = BELOW_FILL
            if r["laps_led"] > 0:
                st["bold"] = True
            if r["led_most"] == 1:
                st["bold"], st["color"] = True, RED
            if r["pole"] == 1:
                st["italic"] = True
            if r["dnf"] == 1:
                st["strike"] = True
            if "fl_race" in r.index and r["fl_race"] == 1:
                st["underline"] = True
            s1, s2 = r["s1_win"] == 1, r["s2_win"] == 1
            if s1 or s2:
                st["bold"] = True
                st["color"] = PINK if (s1 and s2) else BLUE
            if st:
                styles[(i, c)] = st
    return styles


# ----------------------------------------------------------------------------- images
def _fmt(v):
    if v is None or (isinstance(v, float) and pd.isna(v)) or v is pd.NA:
        return ""
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v)


def _dur(m):
    if m is None:
        return ""
    m = int(round(m))
    return f"{m // 60}h {m % 60:02d}m" if m >= 60 else f"{m} min"


def _draw_bar(ax, W, top, bar):
    """Race-summary bar: green flag laps, yellow cautions, red flag marker, stage boundaries."""
    from matplotlib.patches import Polygon, Rectangle
    total = bar["total"]
    x0, x1 = 0.25, W - 0.25
    span = x1 - x0
    X = lambda lap: x0 + span * min(max(lap, 0), total) / total
    by, bh = top + 0.5, 0.26
    ax.add_patch(Rectangle((x0, by), span, bh, color="#9bd9a8", lw=0))
    for a, b in bar["cautions"]:
        ax.add_patch(Rectangle((X(a - 1), by), max(X(b) - X(a - 1), 0.025), bh, color="#ffd93d", lw=0))
    for lap, mins in bar["reds"]:                      # a point in the race, not a span of laps
        x = X(lap)
        ax.plot([x, x], [by - 0.04, by + bh + 0.04], color="#d11a1a", lw=2.6, solid_capstyle="butt")
        ax.add_patch(Polygon([(x - 0.07, by - 0.13), (x + 0.07, by - 0.13), (x, by - 0.03)],
                             closed=True, color="#d11a1a", lw=0))
        label = "Red flag" + (f" · {_dur(mins)}" if mins is not None else "")
        ax.text(x, by - 0.16, label, ha="center", va="bottom", fontsize=8,
                color="#d11a1a", fontweight="bold")
    edges = [0.0] + bar["stages"] + [total]
    for i, (a, b) in enumerate(zip(edges[:-1], edges[1:])):
        label = f"Stage {i + 1}" if i < len(edges) - 2 else ("Final" if len(edges) > 2 else "")
        ax.text((X(a) + X(b)) / 2, by + bh + 0.13, label, ha="center", va="center", fontsize=8, color="#333333")
    for e in bar["stages"]:
        ax.plot([X(e)] * 2, [by - 0.05, by + bh + 0.05], color="black", lw=1.4)
    ax.add_patch(Rectangle((x0, by), span, bh, fill=False, ec="black", lw=0.8))
    claps = int(sum(b - a + 1 for a, b in bar["cautions"]))
    parts = [f"{len(bar['cautions'])} cautions ({claps} laps)"] if bar["cautions"] else []
    if bar["reds"]:
        tot = sum(m for _, m in bar["reds"] if m is not None)
        parts.append(f"{len(bar['reds'])} red flag" + ("s" if len(bar["reds"]) > 1 else "")
                     + (f" ({_dur(tot)})" if tot else ""))
    parts.append(f"{int(total)} laps")
    ax.text(x0, by + bh + 0.3, "Green = racing · Yellow = caution · Red = stoppage     " + " · ".join(parts),
            ha="left", va="center", fontsize=8, style="italic", color="#444444")


def render_png(df, title, path, n_highlight=0, blue_rows=(), footer="", styles=None, legend=None, row_fills=None, bar=None):
    """Draw a DataFrame as a phone-friendly table image."""
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    from matplotlib.patches import Rectangle

    cols = list(df.columns)
    cells = [[_fmt(v) for v in row] for row in df.itertuples(index=False)]
    char, pad, rh = 0.078, 0.2, 0.26
    widths = [max([len(c)] + [len(r[i]) for r in cells]) * char + pad for i, c in enumerate(cols)]
    W = sum(widths)
    title_h = 0.8 if title else 0.1
    n_foot = len(footer.split("\n")) if footer else 0
    foot_h = (0.12 + 0.22 * n_foot if footer else 0.1) + (0.35 if legend else 0)
    bar_h = 1.25 if bar else 0
    H = title_h + bar_h + rh * (len(cells) + 1) + foot_h
    fig = plt.figure(figsize=(W, H))
    ax = fig.add_axes([0, 0, 1, 1])
    ax.set_xlim(0, W)
    ax.set_ylim(H, 0)
    ax.axis("off")

    if title:
        ax.text(W / 2, title_h / 2, title, ha="center", va="center", fontsize=13,
                fontweight="bold", color="#0a9a4a", linespacing=1.3)

    xs = [sum(widths[:i]) for i in range(len(widths))]
    y0 = title_h + bar_h
    if bar:
        _draw_bar(ax, W, title_h, bar)
    if row_fills is not None:
        for i, col in enumerate(row_fills):
            if col:
                ax.add_patch(Rectangle((0, y0 + rh * (i + 1)), W, rh, color=col, lw=0))
    else:
        for i in range(min(n_highlight, len(cells))):
            ax.add_patch(Rectangle((0, y0 + rh * (i + 1)), W, rh, color="#fff6a0", lw=0))
    for (i, cname), st in (styles or {}).items():
        if st.get("fill") and cname in cols and i < len(cells):
            j = cols.index(cname)
            ax.add_patch(Rectangle((xs[j], y0 + rh * (i + 1)), widths[j], rh, color=st["fill"], lw=0))
    for i in range(len(cells) + 1):
        ax.plot([0, W], [y0 + rh * (i + 1)] * 2, color="#dddddd", lw=0.5)
    if 0 < n_highlight < len(cells):
        ax.plot([0, W], [y0 + rh * (n_highlight + 1)] * 2, color="black", lw=1.2, ls=(0, (5, 3)))
    ax.plot([0, W], [y0 + rh] * 2, color="black", lw=1.2)
    ax.plot([0, W], [y0] * 2, color="black", lw=1.2)

    for j, c in enumerate(cols):
        left = c == "Driver"
        ax.text(xs[j] + (0.08 if left else widths[j] / 2), y0 + rh / 2, c,
                ha="left" if left else "center", va="center", fontsize=9.5, fontweight="bold")
    for i, row in enumerate(cells):
        y = y0 + rh * (i + 1) + rh / 2
        for j, val in enumerate(row):
            left = cols[j] == "Driver"
            color, weight = "black", "normal"
            if cols[j] == "+/-" and val.startswith("+"):
                color, weight = "#14902e", "bold"
            elif cols[j] == "+/-" and val.startswith("-"):
                color, weight = "#d11a1a", "bold"
            if left and i in blue_rows:
                color, weight = "#1f5fbf", "bold"
            fstyle, strike, uline = "normal", False, False
            st = (styles or {}).get((i, cols[j]))
            if st:
                color = st.get("color", color)
                weight = "bold" if st.get("bold") else weight
                fstyle = "italic" if st.get("italic") else "normal"
                strike = st.get("strike", False)
                uline = st.get("underline", False)
            cx = xs[j] + (0.08 if left else widths[j] / 2)
            ax.text(cx, y, val, ha="left" if left else "center", va="center", fontsize=9.5,
                    color=color, fontweight=weight, fontstyle=fstyle)
            if strike and val:
                hw = len(val) * 0.042
                ax.plot([cx - hw, cx + hw], [y, y], color=color, lw=1.1)
            if uline and val:
                hw = len(val) * 0.042
                ax.plot([cx - hw, cx + hw], [y + 0.085, y + 0.085], color=color, lw=1.1)
    if footer:
        for k, line in enumerate(footer.split("\n")):
            ax.text(0.1, H - foot_h + 0.17 + 0.22 * k, line, ha="left", va="center",
                    fontsize=8.5, style="italic", color="#444444")
    if legend:
        x, ly = 0.1, H - 0.2
        for label, st in legend:
            ax.text(x, ly, label, ha="left", va="center", fontsize=9,
                    color=st.get("color", "black"), fontweight="bold" if st.get("bold") else "normal",
                    fontstyle="italic" if st.get("italic") else "normal")
            w = len(label) * 0.08
            if st.get("strike"):
                ax.plot([x, x + w], [ly, ly], color="black", lw=1.1)
            if st.get("underline"):
                ax.plot([x, x + w], [ly + 0.09, ly + 0.09], color="black", lw=1.1)
            x += w + 0.25
    fig.savefig(path, dpi=200, facecolor="white")
    plt.close(fig)


def png_standings(t2, title, path, n_races, size, grouped, results=None, shade=True, shade_below=False):
    sdf = t2.copy()
    blue = [i for i, w in enumerate(sdf["Wins"]) if w > 0]
    last = [c for c in sdf.columns if c.startswith("R") and c[1:].isdigit()][-n_races:] if n_races else []
    cols = ["Rank", "+/-", "Car", "Driver", "Points", "Behind", "Wins", "Stage Pts",
            "Top 5", "Top 10", "DNFs", "Avg Finish"]
    if "Best Lap Pts" in t2:
        cols += ["Best Lap Pts"]
    if t2["Rating"].notna().any():
        cols += ["Rating", "Last 5 Rtg"]
    cols += last
    sdf = sdf[cols]
    styles = cell_styles(results, t2, last, shade, shade_below) if results is not None else None
    sdf["+/-"] = sdf["+/-"].map(lambda v: f"+{v}" if v > 0 else (str(v) if v < 0 else ""))
    sdf["Behind"] = sdf["Behind"].map(lambda v: "-" if v == 0 else str(v))
    sdf["Avg Finish"] = sdf["Avg Finish"].map(lambda v: f"{v:.1f}")
    for c in ("Rating", "Last 5 Rtg"):
        if c in sdf:
            sdf[c] = sdf[c].map(lambda v: "" if pd.isna(v) else f"{v:.1f}")
    sdf = sdf.rename(columns={"Stage Pts": "Stg Pts", "Avg Finish": "Avg Fin", "Top 5": "T5", "Top 10": "T10",
                              "Rating": "Rtg", "Last 5 Rtg": "L5 Rtg",
                              "Best Lap Pts": "BL"})
    for c in ["Wins", "T5", "T10", "DNFs", "BL"]:
        if c not in sdf:
            continue
        sdf[c] = sdf[c].map(lambda v: "·" if v == 0 else v)
    foot = ("Yellow = in the Chase field · Blue = race winner · unofficial standings"
            + (" · Finish shading: dark green top 5, light green 6-10" + (", grey 16+" if shade_below else "")
               if shade else ""))
    render_png(sdf, title, path, n_highlight=size, blue_rows=blue, footer=foot,
               styles=styles,
               legend=(LEGEND if (results is not None and "fl_race" in results) else LEGEND[:-1]) if styles else None)


TOP5_FILL, TOP10_FILL, BELOW_FILL = "#b7e4b7", "#e2f4e2", "#e6e6e6"


def png_race(t1, title, path, bar=None):
    """Finish-tier shading: top 5 and 6-10 shaded, 11-15 left white, 16+ shaded grey."""
    fills = [TOP5_FILL if f <= 5 else TOP10_FILL if f <= 10 else None if f <= 15 else BELOW_FILL
             for f in t1["Finish"]]
    rdf = t1.drop(columns=["DNF"]) if "DNF" in t1 else t1.copy()
    for c in ("Rating", "Last 5 Rtg"):
        if c in rdf:
            rdf[c] = rdf[c].map(lambda v: "" if pd.isna(v) else f"{v:.1f}")
    rdf = rdf.rename(columns={"Laps Led": "Led", "Total Pts": "Total", "Rank After": "Rank",
                              "Rating": "Rtg", "Last 5 Rtg": "L5 Rtg"})
    styles = {}
    if "FL Rk" in rdf:
        styles = {(i, "Best Lap"): dict(bold=True, underline=True)
                  for i, rk in enumerate(rdf["FL Rk"]) if pd.notna(rk) and rk == 1}
        rdf["Best Lap"] = rdf["Best Lap"].map(lambda v: "" if pd.isna(v) else f"{v:.3f}")
        rdf["FL Rk"] = rdf["FL Rk"].map(lambda v: "" if pd.isna(v) else str(int(v)))
        if "Mph" in rdf:
            rdf["Mph"] = rdf["Mph"].map(lambda v: "" if pd.isna(v) else f"{v:.1f}")
    render_png(rdf, title, path, row_fills=fills, blue_rows=[0] if len(rdf) else [], styles=styles, bar=bar,
               footer="Dark green = top 5 · light green = 6-10 · white = 11-15 · grey = 16+\n"
                      "Blue name = winner · Rtg = driver rating · L5 = last 5 rated starts"
                      + (" · Underline = fastest lap" if "FL Rk" in rdf else ""))


# ----------------------------------------------------------------------------- export
def _style_excel(ws, t2, styles):
    from openpyxl.styles import Font, PatternFill
    for (i, col), st in styles.items():
        cell = ws.cell(row=i + 2, column=list(t2.columns).index(col) + 1)
        if st.get("fill"):
            cell.fill = PatternFill("solid", start_color=st["fill"].lstrip("#"), end_color=st["fill"].lstrip("#"))
        cell.font = Font(name="Arial", bold=st.get("bold", False), italic=st.get("italic", False),
                         strike=st.get("strike", False),
                         underline="single" if st.get("underline") else None,
                         color=(st["color"].lstrip("#") if "color" in st else None))
    ws.cell(row=len(t2) + 3, column=1).value = (
        "Key: italic = pole · bold = led a lap · red = led most laps · blue = won a stage · pink = won both stages · "
        "strikethrough = DNF")


def export_week(results, adj, week, outdir, chase=None, xlsx=True, png=True,
                png_races=5, title="NASCAR Cup Series", size=16, events=None,
                shade=True, shade_below=False):
    os.makedirs(outdir, exist_ok=True)
    t1 = race_table(results, adj, week, chase)
    t2 = standings_table(results, adj, week, chase)
    rows = results[results.race_num == week]
    name = rows["race_name"].iloc[0] if len(rows) else ""
    track = rows["track"].iloc[0] if len(rows) else ""
    label = track or name
    stem = f"week_{week:02d}"
    t1.to_csv(os.path.join(outdir, f"{stem}_race.csv"), index=False)
    t2.to_csv(os.path.join(outdir, f"{stem}_standings.csv"), index=False)
    if xlsx:
        with pd.ExcelWriter(os.path.join(outdir, f"{stem}.xlsx"), engine="openpyxl") as xw:
            t1.to_excel(xw, sheet_name=f"Race {week}"[:31], index=False)
            t2.to_excel(xw, sheet_name="Season to date", index=False)
            for sheet, extra in (("Fastest laps", fastest_table(results, week)),
                                 ("Race events", events_table(events, week))):
                if extra is not None:
                    extra.to_excel(xw, sheet_name=sheet, index=False)
            _style_excel(xw.sheets["Season to date"], t2, cell_styles(
                results, t2, [c for c in t2.columns if c.startswith("R") and c[1:].isdigit()],
                shade, shade_below))
    if png:
        png_standings(t2, f"{title}\nafter {label}", os.path.join(outdir, f"{stem}_standings.png"),
                      png_races, size, _grouped(week, chase), results=results,
                      shade=shade, shade_below=shade_below)
        png_race(t1, f"{title}\nRace {week} results: {name or label}",
                 os.path.join(outdir, f"{stem}_race.png"), bar=bar_for_race(events, week))
    return label


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("results")
    ap.add_argument("--adjustments")
    ap.add_argument("--no-cell-shading", action="store_true",
                    help="no finish-position shading in the race-by-race cells of the points view")
    ap.add_argument("--shade-below", action="store_true",
                    help="also shade finishes of 16th or worse grey in the points view")
    ap.add_argument("--no-auto", action="store_true",
                    help="don't pick up fastest_laps.csv / race_events.csv automatically")
    ap.add_argument("--events", help="race_events.csv: stage ends, cautions, red flags for the race bar")
    ap.add_argument("--fastest-laps", help="CSV: race_num, driver, best_lap_time [, best_lap_speed, best_lap_num]")
    ap.add_argument("--week", type=int)
    ap.add_argument("--all-weeks", action="store_true")
    ap.add_argument("--outdir", default="out")
    ap.add_argument("--no-xlsx", action="store_true")
    ap.add_argument("--no-png", action="store_true")
    ap.add_argument("--png-races", type=int, default=5, help="recent race finishes shown in the image")
    ap.add_argument("--title", default="NASCAR Cup Series")
    ap.add_argument("--reset-after", type=int, default=26)
    ap.add_argument("--chase-size", type=int, default=16)
    ap.add_argument("--reset-base", type=int, default=2000)
    ap.add_argument("--no-reset", action="store_true")
    a = ap.parse_args()

    results = load_results(a.results)
    adj = load_adjustments(a.adjustments)
    here = os.path.dirname(os.path.abspath(a.results))
    if not a.no_auto:                         # pick up the side files sitting next to the results file
        if not a.fastest_laps and os.path.exists(os.path.join(here, "fastest_laps.csv")):
            a.fastest_laps = os.path.join(here, "fastest_laps.csv")
            print("using fastest_laps.csv found next to the results file")
        if not a.events and os.path.exists(os.path.join(here, "race_events.csv")):
            a.events = os.path.join(here, "race_events.csv")
            print("using race_events.csv found next to the results file")
    if a.fastest_laps:
        results = merge_fastest(results, a.fastest_laps)
    events = load_events(a.events) if a.events else None
    chase = None
    if not a.no_reset:
        adj, chase = make_chase(results, adj, a.reset_after, a.chase_size, a.reset_base)
    weeks = sorted(results.race_num.unique()) if a.all_weeks else [a.week or int(results.race_num.max())]
    for w in weeks:
        label = export_week(results, adj, int(w), a.outdir, chase, xlsx=not a.no_xlsx,
                            png=not a.no_png, png_races=a.png_races, title=a.title, size=a.chase_size,
                            events=events, shade=not a.no_cell_shading, shade_below=a.shade_below)
        print(f"week {w} {label}: written to {a.outdir}/")


if __name__ == "__main__":
    main()
